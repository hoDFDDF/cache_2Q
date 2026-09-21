"""Integer Zipf requests -> LFU executable -> Excel table and plots."""

import argparse
import csv
import heapq
import math
import subprocess
from datetime import datetime
from pathlib import Path

import matplotlib

matplotlib.use("Agg")  # Save plots without requiring a graphical session.
import matplotlib.pyplot as plt
import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill


PROJECT_ROOT = Path(__file__).resolve().parents[2]


def parseArgs():
    parser = argparse.ArgumentParser(description="LFU benchmark: Zipf requests, tables and plots")
    parser.add_argument("--cache-bin", type=Path, default=PROJECT_ROOT / "build/modules/LFU-cache")
    parser.add_argument("--requests", type=int, default=100_000)
    parser.add_argument("--keys", type=int, default=10_000)
    parser.add_argument("--capacities", type=int, nargs="+", default=[100, 500, 1000])
    parser.add_argument("--alphas", type=float, nargs="+", default=[0.8, 1.0, 1.2])
    parser.add_argument("--seed", type=int, default=42)
    parser.add_argument("--output", type=Path, default=PROJECT_ROOT / "build/lfu_python_results")
    args = parser.parse_args()
    if not 0 <= args.requests <= 2**31 - 1 or not 1 <= args.keys <= 2**31 - 1:
        parser.error("requests must be 0..INT_MAX and keys must be 1..INT_MAX")
    if args.seed < 0 or any(not 0 <= capacity <= 2**31 - 1 for capacity in args.capacities):
        parser.error("seed must be >= 0 and capacities must be 0..INT_MAX")
    if any(not math.isfinite(alpha) or alpha <= 0 for alpha in args.alphas):
        parser.error("alpha must be finite and > 0")
    args.cache_bin = args.cache_bin.resolve()
    if not args.cache_bin.is_file():
        parser.error("Build LFU-cache first: cmake -S . -B build && cmake --build build --target LFU-cache")
    args.capacities = sorted(set(args.capacities))
    args.alphas = sorted(set(args.alphas))
    return args


def generateRequests(keys_count, requests_count, alpha, seed):
    keys = np.arange(1, keys_count + 1, dtype=np.int64)
    weights = keys.astype(float) ** (-alpha)
    probabilities = weights / weights.sum()
    rng = np.random.default_rng(seed)
    requests = rng.choice(keys, size=requests_count, p=probabilities)
    return requests, probabilities


def countReferenceHits(requests, capacity):
    """Independent LFU + LRU: dictionary and a heap, without cache list iterators."""
    entries = {}
    priorities = []
    hits = 0
    for tick, key in enumerate(requests):
        if capacity == 0:
            continue
        if key in entries:
            frequency = entries[key][0] + 1
            hits += 1
        else:
            frequency = 1
            if len(entries) == capacity:
                while priorities:
                    old_frequency, old_tick, old_key = heapq.heappop(priorities)
                    # Hits leave old heap records; ignore those outdated records.
                    if entries.get(old_key) == (old_frequency, old_tick):
                        del entries[old_key]
                        break
        entries[key] = (frequency, tick)
        heapq.heappush(priorities, (frequency, tick, key))
    return hits


def runCache(cache_bin, input_path, requests_count):
    with input_path.open() as input_file:
        result = subprocess.run(
            [str(cache_bin)], stdin=input_file, capture_output=True, text=True, timeout=30
        )
    input_path.with_suffix(".out").write_text(result.stdout, encoding="utf-8")
    input_path.with_suffix(".err").write_text(result.stderr, encoding="utf-8")
    if result.returncode != 0:
        raise RuntimeError(f"LFU-cache exited with {result.returncode}: {result.stderr.strip()}")
    try:
        hits = int(result.stdout.strip())
    except ValueError as error:
        raise RuntimeError(f"LFU-cache did not return one integer: {result.stdout!r}") from error
    if not 0 <= hits <= requests_count:
        raise RuntimeError(f"Invalid hit count: {hits}")
    return hits


def saveTable(rows, output_dir):
    headers = list(rows[0])
    with (output_dir / "results.csv").open("w", newline="", encoding="utf-8-sig") as output:
        writer = csv.DictWriter(output, fieldnames=headers, delimiter=";")
        writer.writeheader()
        writer.writerows(rows)

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "LFU results"
    sheet.append(headers)
    for row in rows:
        sheet.append(list(row.values()))
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="244062")
        cell.font = Font(color="FFFFFF", bold=True)
    for column in sheet.columns:
        sheet.column_dimensions[column[0].column_letter].width = 18
        if column[0].value == "hit_rate":
            for cell in column[1:]:
                cell.number_format = "0.00%"
    workbook.save(output_dir / "results.xlsx")


def savePlots(rows, streams, output_dir):
    plt.style.use("seaborn-v0_8-whitegrid")
    figure, (hits_axis, zipf_axis) = plt.subplots(1, 2, figsize=(13, 5), layout="constrained")
    for alpha, probabilities, frequencies in streams:
        selected = [row for row in rows if row["alpha"] == alpha]
        hits_axis.plot(
            [row["capacity"] for row in selected],
            [100 * row["hit_rate"] for row in selected],
            marker="o", linewidth=2, label=f"alpha = {alpha:g}",
        )
        # Show the first 100 ranks, so the popularity plot stays readable.
        count = min(100, len(probabilities))
        ranks = np.arange(1, count + 1)
        line, = zipf_axis.loglog(ranks, probabilities[:count], label=f"alpha = {alpha:g}")
        observed = frequencies[:count]
        mask = observed > 0
        zipf_axis.scatter(ranks[mask], observed[mask], color=line.get_color(), s=12, alpha=0.7)

    hits_axis.set(title="LFU: hit rate (cold start)", xlabel="Cache capacity (keys)",
                  ylabel="Hit rate (%)", ylim=(0, 100))
    zipf_axis.set(title="Zipf: lines = theory, dots = generated requests",
                  xlabel="Key / popularity rank", ylabel="Request probability")
    for axis in (hits_axis, zipf_axis):
        axis.legend()
    figure.savefig(output_dir / "plots.png", dpi=180)
    plt.close(figure)


def main():
    args = parseArgs()
    output_dir = args.output.resolve() / datetime.now().strftime("run-%Y%m%d-%H%M%S-%f")
    output_dir.mkdir(parents=True)
    inputs_dir = output_dir / "inputs"
    inputs_dir.mkdir()
    print(f"Results: {output_dir}", flush=True)
    print(f"{'alpha':>7} {'capacity':>10} {'hits':>10} {'misses':>10} {'hit rate':>10} {'status':>8}")
    rows = []
    streams = []
    for alpha_index, alpha in enumerate(args.alphas):
        requests, probabilities = generateRequests(args.keys, args.requests, alpha, args.seed)
        frequencies = np.bincount(requests, minlength=args.keys + 1)[1:] / max(1, args.requests)
        streams.append((alpha, probabilities, frequencies))
        requests_text = "\n".join(map(str, requests)) + "\n"
        for capacity in args.capacities:
            input_path = inputs_dir / f"alpha{alpha_index}_capacity{capacity}.in"
            input_path.write_text(f"{args.requests} {capacity}\n" + requests_text, encoding="utf-8")
            hits = runCache(args.cache_bin, input_path, args.requests)
            expected_hits = countReferenceHits(requests, capacity)
            row = dict(alpha=alpha, capacity=capacity, requests=args.requests, keys=args.keys,
                       seed=args.seed, hits=hits, misses=args.requests - hits,
                       hit_rate=hits / max(1, args.requests), reference_hits=expected_hits,
                       status="PASS" if hits == expected_hits else "FAIL")
            rows.append(row)
            print(f"{alpha:7g} {capacity:10d} {hits:10d} {row['misses']:10d} "
                  f"{row['hit_rate']:10.2%} {row['status']:>8}", flush=True)
    saveTable(rows, output_dir)
    savePlots(rows, streams, output_dir)
    print("Saved: results.xlsx, results.csv, plots.png and inputs/")
    return 0 if all(row["status"] == "PASS" for row in rows) else 1


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except (OSError, RuntimeError, subprocess.TimeoutExpired) as error:
        raise SystemExit(f"Benchmark error: {error}") from error
